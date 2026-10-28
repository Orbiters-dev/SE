# -*- coding: utf-8 -*-
"""새 계약서 양식(수기 서식) → 생성기용 앵커 사본 생성.
- {{VAR}} = generate_influencer_contract.py 치환용
- {{名前...}} = DocuSeal 필드 태그 (DOCX 업로드 시 자동 필드화)
세은 원본은 건드리지 않고 앵커 사본만 출력 폴더에 저장.
세은이 원본 양식을 개정하면 이 스크립트를 재실행해 앵커 사본을 갱신한다.

7/31 렌더 실측 교훈 (전부 DocuSeal 서명 페이지 실측으로 확정):
- 태그 문자열 길이 = 필드 폭·배경 → role=First Party/type=text/required=true 는 생략 (기본값).
  장문 태그는 라벨을 덮고 서명란을 다음 페이지로 밀어냄.
- 원본의 페이지 브레이크(빈 단락 br + 署名欄 pageBreakBefore) 제거 — 빈 페이지 금지 (세은).
- 希望時間帯 = 라벨 옆 select 드롭다운 1필드 (체크박스 8개·그리드·별도줄 전부 폐기, 세은 확정).
  DocuSeal에 time picker 타입은 없음. 태그 run 7pt + LEFT 정렬이면 라벨 잠식 없음.
- 配達希望日 = type=date 달력 단독 (指定なし 체크박스 제거, ※주석이 기입불요 안내).
- 플랫폼 3종 병기 금지 (세은) → {{PLATFORMS}} 앵커, 계약별 합의 플랫폼 기재.
"""
import os
import re
import sys
from pathlib import Path
from docx import Document

try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

# 세은 원본 양식 폴더 / 앵커 사본 출력 폴더 — 다른 PC(핸드오버)에서는 .env로 지정
NEW_DIR = Path(os.getenv(
    "CONTRACT_SOURCE_DIR",
    r"C:\Users\orbit\Desktop\s\contents\instagram\인플루언서 계약서(서명 無 ver)"))
OUT_DIR = Path(os.getenv(
    "CONTRACT_TEMPLATE_DIR",
    r"C:\Users\orbit\Desktop\s\참고자료\인플루언서 컨택\형식"))
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def tag(name, ftype=None, required=True, options=None):
    parts = [name]
    if ftype:
        parts.append("type=" + ftype)
    if options:
        parts.append("options=" + options)
    if not required:
        parts.append("required=false")
    return "{{" + ";".join(parts) + "}}"


# 본문 치환 (유상/무상 각각 — 원문 표기가 미묘하게 다름)
BODY_PAID = [
    ("〇〇年〇〇月〇〇日付で", "{{AGREEMENT_DATE}}付で"),
    ("@ IDとして知られる 氏名", "@{{ACCOUNT_HANDLE}} として知られる {{INFLUENCER_NAME}}"),
    ("動画コンテンツ 〇本", "動画コンテンツ {{VIDEO_COUNT}}本"),
    ("（Instagram Reels/YouTube shorts/TikTok）", "（{{PLATFORMS}}）"),
    ("商品名：〇〇〇〇", "商品名：{{PRODUCT_NAME}}"),
    ("数量：〇点", "数量：{{QUANTITY}}点"),
    ("カラー・サイズ等：〇〇/〇〇ml", "カラー・サイズ等：{{COLOR_SIZE}}"),
    ("¥〇〇〇〇（税込み）", "¥{{PAYMENT_AMOUNT}}（税込み）"),
    ("広告表記（#grosmimi #グロミミ #ストローマグ #スマートマグ）", "広告表記（{{HASHTAGS}}）"),
]
BODY_GIFT = [
    ("〇〇〇〇年〇〇月〇〇日付で", "{{AGREEMENT_DATE}}付で"),
    ("@ ID として知られる 氏名", "@{{ACCOUNT_HANDLE}} として知られる {{INFLUENCER_NAME}}"),
    ("動画コンテンツ 〇本", "動画コンテンツ {{VIDEO_COUNT}}本"),
    ("（Instagram Reels/YouTube shorts/TikTok）", "（{{PLATFORMS}}）"),
    ("商品名：〇〇〇〇", "商品名：{{PRODUCT_NAME}}"),
    ("数量：〇点", "数量：{{QUANTITY}}点"),
    ("カラー・サイズ等：〇〇/〇〇ml", "カラー・サイズ等：{{COLOR_SIZE}}"),
    ("広告表記（例：#grosmimi #グロミミ #ストローマグ #スマートマグ）", "広告表記（{{HASHTAGS}}）"),
]

U = r"[_＿]+"  # 밑줄 공란 (반각/전각)


def sub_label(text, label, replacement):
    pat = re.compile(re.escape(label) + r"：\s*" + U)
    return pat.sub(label + "：" + replacement, text)


def transform_line(text, mode):
    orig = text
    if mode == "ko":  # 【甲】 서명란 — 회사측 자동 기입
        if re.match(r"^氏名：\s*" + U + r"\s*$", text):
            return "氏名：{{COMPANY_NAME}}"
        if re.match(r"^署名：\s*" + U + r"\s*$", text):
            return "署名：{{COMPANY_NAME}}"
        if re.match(r"^日付：\s*" + U + r"年" + U + r"月" + U + r"日\s*$", text):
            return "日付：{{AGREEMENT_DATE}}"
    elif mode == "otsu":  # 【乙】 인플루언서 — DocuSeal 입력 필드
        if re.match(r"^氏名：\s*" + U, text):
            return "氏名：" + tag("氏名")
        if re.match(r"^メールアドレス：\s*" + U, text):
            return "メールアドレス：" + tag("メールアドレス")
        if re.match(r"^Instagramハンドル：@\s*" + U, text):
            return "Instagramハンドル：@" + tag("Instagram")
        if re.match(r"^署名：\s*" + U, text):
            return "署名：" + tag("署名", "signature")
        if re.match(r"^日付：\s*" + U, text):
            return "日付：" + tag("日付", "date")
    elif mode == "addr":  # 【商品配送先住所】
        for label in ("郵便番号", "住所", "宛名", "電話番号"):
            if re.match(r"^" + re.escape(label) + r"：\s*" + U, text):
                # 住所는 긴 입력값 대비 설명형 필드명으로 폭 확보 (필드명 = 태그 폭)
                fname = "住所（番地・建物名まで）" if label == "住所" else label
                return sub_label(text, label, tag(fname))
        if text.startswith("配達希望日："):
            # 날짜는 달력 선택 (type=date). 指定なし 체크박스는 제거 — ※주석이 기입불요 안내
            return "配達希望日：" + tag("配達希望日", "date", required=False)
        if text.startswith("希望時間帯："):
            # 최종 확정: 라벨 바로 옆 select 드롭다운 1필드 (선택지 리스트·체크박스 없음)
            opts = "指定なし,午前中,12〜14時,14〜16時,16〜18時,18〜20時,19〜21時,20〜21時"
            return "希望時間帯：" + tag("希望時間帯", "select", required=False, options=opts)
        if text.startswith("※お届けは"):
            # 세은 7/31 지정 문구 추가 (희망 일시 없으면 기입 불요)
            return text + "配達希望日がない場合は、記入不要です。"
    elif mode == "bank":  # 【報酬振込先】 (유상만)
        for label in ("金融機関名", "支店名（店番号）", "口座種別", "口座番号", "口座名義（カナ）"):
            text = sub_label(text, label, tag(label.split("（")[0]))
        return text
    return orig


def remove_page_breaks(doc):
    """명시적 페이지 브레이크 제거 (빈 페이지 금지 — 세은 7/31 지시)"""
    removed = 0
    for para in doc.paragraphs:
        for br in para._p.findall(f".//{{{W_NS}}}br"):
            if br.get(f"{{{W_NS}}}type") == "page":
                br.getparent().remove(br)
                removed += 1
        for pbb in para._p.findall(f".//{{{W_NS}}}pageBreakBefore"):
            pbb.getparent().remove(pbb)
            removed += 1
    return removed


def process(src, dst, body_pairs):
    if not src.exists():
        print(f"[ERROR] 원본 양식 없음: {src}")
        print("  → .env CONTRACT_SOURCE_DIR 로 세은 원본 양식 폴더를 지정하세요.")
        sys.exit(1)
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(src))
    n_br = remove_page_breaks(doc)
    mode = "body"
    body_hits = {old: 0 for old, _ in body_pairs}
    changed = 0
    for para in doc.paragraphs:
        t = para.text
        # 단락 안에 <w:br/> 줄바꿈으로 여러 줄이 붙어 있을 수 있음 → 줄 단위 처리
        out_lines = []
        for line in t.split("\n"):
            if "署名欄" in line:
                mode = "ko"
            elif line.startswith("【乙】"):
                mode = "otsu"
            elif line.startswith("【商品配送先住所】"):
                mode = "addr"
            elif line.startswith("【報酬振込先】"):
                mode = "bank"

            if mode == "body":
                for old, rep in body_pairs:
                    if old in line:
                        line = line.replace(old, rep)
                        body_hits[old] += 1
            else:
                line = transform_line(line, mode)
            out_lines.append(line)
        new = "\n".join(out_lines)

        if new != t:
            changed += 1
            if para.runs:
                for r in para.runs:
                    r.text = ""
                if "type=checkbox" in new or "type=date" in new or "type=select" in new:
                    # 태그 조각만 폰트 축소 (태그 bbox = 필드 폭·배경 → 줄 깨짐·과대 박스 방지)
                    # 라벨·구분자는 원 크기 유지. justify가 태그를 늘리지 않게 왼쪽 정렬
                    from docx.shared import Pt
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    pieces = re.split(r"(\{\{[^}]+\}\})", new)
                    first = True
                    for piece in pieces:
                        if not piece:
                            continue
                        if first:
                            para.runs[0].text = piece
                            run = para.runs[0]
                            first = False
                        else:
                            run = para.add_run(piece)
                        if piece.startswith("{{"):
                            # options 장문 select는 7pt (밴드 얇게), 그 외 8pt
                            run.font.size = Pt(7 if "type=select" in piece else 8)
                else:
                    para.runs[0].text = new
            else:
                para.add_run(new)

    doc.save(str(dst))
    print(f"{dst.name}: {changed} 단락 치환, 페이지 브레이크 {n_br}개 제거")
    for old, n in body_hits.items():
        status = "OK" if n > 0 else "★MISS★"
        print(f"  [{status}] {old} x{n}")


if __name__ == "__main__":
    process(NEW_DIR / "newインフルエンサーコンテンツ契約書（有償）.docx",
            OUT_DIR / "newインフルエンサーコンテンツ契約書（有償）_anchored.docx", BODY_PAID)
    process(NEW_DIR / "newインフルエンサーコンテンツ契約書（無償）.docx",
            OUT_DIR / "newインフルエンサーコンテンツ契約書（無償）_anchored.docx", BODY_GIFT)

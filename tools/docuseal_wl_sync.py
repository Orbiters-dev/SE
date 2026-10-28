"""
docuseal_wl_sync.py — DocuSeal 유상 계약 → WL Code & Payment 탭 자동 동기화
============================================================================
KPI 보드 예산(Paid Fee) 집계의 구멍(시트 Deal Type/Fee 공란) 방지.

흐름:
  1. DocuSeal /api/submissions 전체 조회 (GROSMIMI_JP 계약만)
  2. 서명 완료(completed) 유상 건의 핸들·보수 확보
     - 1순위: external_id 구조화 포맷 (h-/t-/f- 세그먼트, 신규 발송분)
     - 2순위: .tmp/docuseal_submissions/submission_{id}.json → 계약 DOCX 파싱
       (핸들 = 본문 @id, 보수 = 「報酬額： ¥NN,NNN」)
     - 3순위: 계약서 폴더에서 서명자 실명으로 DOCX 탐색
  3. WL Code & Payment 탭에서 Creator ID 매칭
  4. Deal Type·Fee(JYP) 둘 다 공란인 행만 채움 (기존 값 절대 덮어쓰기 X)
     - 매칭 행 0개/복수 후보/값 불일치 → 수동 확인 리포트만

입금 여부(Payment Status)는 은행 데이터가 없어 자동화 대상 아님 — 세은 수동 유지.

사용:
  python tools/docuseal_wl_sync.py            # dry-run (변경 미리보기)
  python tools/docuseal_wl_sync.py --apply    # 시트 실제 기입
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import requests
from dotenv import load_dotenv

sys.path.insert(0, str(Path(__file__).parent))
from sheets_utils import get_sheets_client, safe_range  # noqa: E402

load_dotenv()

DOCUSEAL_BASE_URL = os.getenv("DOCUSEAL_BASE_URL", "https://docuseal.orbiters.co.kr")
DOCUSEAL_API_KEY = os.getenv("DOCUSEAL_API_KEY")
API_HEADERS = {"X-Auth-Token": DOCUSEAL_API_KEY}

TRACKER_SHEET_ID = "13S1cST2ukuNNHNUmyXAr1HuaYPsuuK_EfUQ10IWIQsE"
WL_TAB_NAME = "WL Code & Payment"
HEADER_ROW = 3  # 1-based

SUBMISSION_RECORDS = Path(".tmp/docuseal_submissions")
# 계약 DOCX 폴더 — generate_influencer_contract.py 의 CONTRACT_OUTPUT_DIR 와 동일 기본값
CONTRACT_DIR = Path(os.getenv(
    "CONTRACT_OUTPUT_DIR",
    r"C:\Users\orbit\Desktop\s\contents\instagram\인플루언서 계약서(서명 無 ver)"))

# 계약 본문에 등장하는 자사/무관 핸들 — 인플루언서 핸들 후보에서 제외
OWN_HANDLES = {"grosmimi_japan", "grosmimi_usa", "grosmimi.official", "onzenna.official"}


# ── DocuSeal ────────────────────────────────────────────────────────────────

def fetch_all_submissions() -> list[dict]:
    """전체 submission 조회 — API pagination.next 커서 우선, 없으면 max(id) 폴백."""
    out, after, seen = [], None, set()
    while True:
        params = {"limit": 100}
        if after:
            params["after"] = after
        resp = requests.get(f"{DOCUSEAL_BASE_URL}/api/submissions",
                            headers=API_HEADERS, params=params, timeout=30)
        resp.raise_for_status()
        body = resp.json()
        data = [s for s in body.get("data", []) if s.get("id") and s["id"] not in seen]
        if not data:
            break
        out.extend(data)
        seen.update(s["id"] for s in data)
        after = body.get("pagination", {}).get("next") or max(s["id"] for s in data)
        if len(body.get("data", [])) < 100:
            break
    return out


def is_grosmimi_jp(sub: dict) -> bool:
    ext = sub.get("external_id") or ""
    name = sub.get("name") or ""
    sub_ext = ""
    if sub.get("submitters"):
        sub_ext = sub["submitters"][0].get("external_id") or ""
    return (ext.startswith("GROSMIMI_JP") or sub_ext.startswith("GROSMIMI_JP")
            or name.startswith("JP Contract"))


# ── 핸들·보수 추출 ──────────────────────────────────────────────────────────

def parse_external_id(ext: str) -> dict:
    """신규 구조화 포맷: GROSMIMI_JP_{ts}__h-{handle}__t-{type}__f-{fee}"""
    out = {}
    for seg in ext.split("__")[1:]:
        if seg.startswith("h-"):
            out["handle"] = seg[2:].lower()
        elif seg.startswith("t-"):
            out["collab_type"] = seg[2:]
        elif seg.startswith("f-"):
            try:
                out["fee"] = int(seg[2:].replace(",", ""))
            except ValueError:
                pass
    # 보수가 있는데 type 세그먼트가 빠진 경우 = 유상 (fee 는 유상 계약에만 존재)
    if "fee" in out:
        out.setdefault("collab_type", "paid")
    return out


def docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(pp.text for pp in cell.paragraphs)
    return "\n".join(texts)


def parse_contract_docx(path: Path) -> dict:
    """계약 DOCX에서 핸들·보수 추출. 실측 포맷: 「報酬額： ¥24,200（税込み）」"""
    try:
        full = docx_text(path)
    except Exception as e:  # noqa: BLE001 — 손상/비표준 DOCX 는 미확보로 리포트
        print(f"  [WARN] DOCX 파싱 실패: {path.name} ({e})")
        return {}
    out = {}
    # IG 핸들 — 대문자 표기 계약서 대비 case-insensitive 매칭 후 소문자 정규화
    handles = [h.lower() for h in re.findall(r"@([A-Za-z0-9_.]{3,30})", full)
               if h.lower() not in OWN_HANDLES]
    if handles:
        out["handle"] = handles[0]
    m = re.search(r"報酬額\s*[：:]\s*[¥￥]?\s*([\d,]+)", full)
    if m:
        out["fee"] = int(m.group(1).replace(",", ""))
        out["collab_type"] = "paid"
    return out


def resolve_submission(sub: dict) -> dict:
    """submission 1건 → {handle, fee, collab_type, source} (확보 실패 필드는 없음)."""
    info = {}
    ext = sub.get("external_id") or ""
    if sub.get("submitters"):
        ext = ext or (sub["submitters"][0].get("external_id") or "")
    if "__h-" in ext:
        info.update(parse_external_id(ext))
        info["source"] = "external_id"
        if "fee" in info:
            return info

    # 로컬 submission 기록 → DOCX 경로
    rec_path = SUBMISSION_RECORDS / f"submission_{sub['id']}.json"
    docx_path = None
    if rec_path.exists():
        rec = json.loads(rec_path.read_text(encoding="utf-8"))
        info.setdefault("collab_type", rec.get("collab_type", ""))
        p = Path(rec.get("pdf", ""))
        if p.suffix.lower() == ".pdf":
            p = p.with_suffix(".docx")
        if p.exists():
            docx_path = p

    # 실명으로 DOCX 탐색 (레거시 — 기록 없음)
    if docx_path is None and sub.get("submitters"):
        name = (sub["submitters"][0].get("name") or "").strip()
        if name:
            cands = sorted(CONTRACT_DIR.glob(f"{name}_*.docx"),
                           key=lambda q: q.stat().st_mtime, reverse=True)
            if cands:
                docx_path = cands[0]

    if docx_path is not None:
        parsed = parse_contract_docx(docx_path)
        for k, v in parsed.items():
            info.setdefault(k, v)
        # 報酬額 조항 존재 = 유상 계약 확정 (gifting 템플릿엔 보수 조항 없음).
        # 발송 시 --type 기본값(gifting)이 기록에 남은 경우를 DOCX 실측으로 교정.
        if "fee" in parsed:
            info["collab_type"] = "paid"
        info.setdefault("source", f"docx:{docx_path.name}")
    return info


# ── WL 탭 ──────────────────────────────────────────────────────────────────

def norm_handle(s: str) -> str:
    return (s or "").strip().lower().lstrip("@")


def fee_num(s: str) -> int | None:
    """시트 Fee 문자열 → 정수. "24,200.00" → 24200. 숫자 없으면 None."""
    digits = re.sub(r"[^\d]", "", (s or "").split(".")[0])
    return int(digits) if digits else None


def load_wl_tab(ws) -> tuple[dict, list[dict]]:
    values = ws.get_all_values()
    hdr = [h.replace("\n", " ").strip() for h in values[HEADER_ROW - 1]]
    col = {name: i for i, name in enumerate(hdr)}
    for need in ("Creator ID", "Deal Type", "Fee(JYP)", "Status", "Creator name"):
        if need not in col:
            raise SystemExit(f"[FATAL] WL 탭 헤더에 '{need}' 없음 — 구조 변경됨")
    rows = []
    for rn, r in enumerate(values[HEADER_ROW:], start=HEADER_ROW + 1):
        def at(name):
            i = col[name]
            return r[i].strip() if i < len(r) else ""
        if not (at("Creator ID") or at("Creator name")):
            continue
        rows.append({
            "rownum": rn, "handle": norm_handle(at("Creator ID")),
            "name": at("Creator name"), "deal": at("Deal Type"),
            "fee": at("Fee(JYP)"), "status": at("Status"),
        })
    return col, rows


def col_letter(idx0: int) -> str:
    letters = ""
    n = idx0 + 1
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


# ── 메인 ───────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description="DocuSeal 유상계약 → WL 탭 Deal Type/Fee 동기화")
    ap.add_argument("--apply", action="store_true", help="시트 실제 기입 (기본은 dry-run)")
    args = ap.parse_args()

    if not DOCUSEAL_API_KEY:
        raise SystemExit("[FATAL] DOCUSEAL_API_KEY 미설정 (.env)")

    subs = [s for s in fetch_all_submissions() if is_grosmimi_jp(s)]
    completed = [s for s in subs
                 if s.get("submitters") and s["submitters"][0].get("status") == "completed"]
    print(f"DocuSeal GROSMIMI_JP submissions: {len(subs)}건 (서명 완료 {len(completed)}건)")

    # 유상 계약만 (핸들+보수 확보된 것)
    paid_contracts = []
    unresolved = []
    for s in completed:
        info = resolve_submission(s)
        name = s["submitters"][0].get("name", "?")
        done = (s["submitters"][0].get("completed_at") or "")[:10]
        if info.get("collab_type") == "paid" and "fee" in info and "handle" in info:
            paid_contracts.append({"sub_id": s["id"], "name": name, "signed": done, **info})
        elif info.get("collab_type") == "paid" or "fee" in info:
            unresolved.append({"sub_id": s["id"], "name": name, "signed": done, **info})
        # gifting(보수 조항 없음)은 대상 아님 — 공란 = Gifted 취급이 정상

    print(f"유상 계약 (핸들·보수 확보): {len(paid_contracts)}건 / 유상인데 미확보: {len(unresolved)}건\n")

    gc = get_sheets_client()
    ws = gc.open_by_key(TRACKER_SHEET_ID).worksheet(WL_TAB_NAME)
    col, rows = load_wl_tab(ws)
    by_handle: dict[str, list[dict]] = {}
    for r in rows:
        if r["handle"]:
            by_handle.setdefault(r["handle"], []).append(r)

    updates, reports = [], []
    for c in paid_contracts:
        cand_rows = by_handle.get(c["handle"], [])
        if not cand_rows:
            reports.append(f"행 없음      | {c['name']} (@{c['handle']}) ¥{c['fee']:,} 서명 {c['signed']} — WL 탭에 Creator ID 행 자체가 없음")
            continue
        already = [r for r in cand_rows
                   if r["deal"].lower() == "paid" and fee_num(r["fee"]) == c["fee"]]
        if already:
            continue  # 이미 기입 완료
        blanks = [r for r in cand_rows if not r["deal"] and not r["fee"]]
        if len(blanks) == 1:
            updates.append((blanks[0], c))
        elif len(blanks) > 1:
            rows_s = ", ".join(f"r{b['rownum']}" for b in blanks)
            reports.append(f"복수 후보    | {c['name']} (@{c['handle']}) ¥{c['fee']:,} — 공란 행 {rows_s} 중 어느 행인지 특정 불가")
        else:
            cur = "; ".join(f"r{r['rownum']} Deal={r['deal'] or '공란'}/Fee={r['fee'] or '공란'}" for r in cand_rows)
            reports.append(f"값 불일치    | {c['name']} (@{c['handle']}) 계약 ¥{c['fee']:,} vs 시트 [{cur}] — 수동 확인")

    print(f"=== 기입 대상 (Deal Type·Fee 공란 → 자동 채움): {len(updates)}건 ===")
    for r, c in updates:
        print(f"  r{r['rownum']:>3} | {c['name']} (@{c['handle']}) | Status={r['status'] or '-'} | "
              f"→ Deal Type=Paid, Fee(JYP)={c['fee']:,} | 서명 {c['signed']} (sub {c['sub_id']}, {c['source']})")

    if reports:
        print(f"\n=== 수동 확인 필요: {len(reports)}건 ===")
        for line in reports:
            print(f"  {line}")
    if unresolved:
        print(f"\n=== 유상인데 핸들/보수 추출 실패: {len(unresolved)}건 ===")
        for u in unresolved:
            print(f"  sub {u['sub_id']} | {u['name']} | 확보: { {k: v for k, v in u.items() if k in ('handle', 'fee')} }")

    if not updates:
        print("\n기입할 변경 없음.")
        return

    if not args.apply:
        print("\n[DRY-RUN] 실제 기입은 --apply 로 실행.")
        return

    deal_col = col_letter(col["Deal Type"])
    fee_col = col_letter(col["Fee(JYP)"])
    data = []
    for r, c in updates:
        data.append({"range": safe_range(WL_TAB_NAME, f"{deal_col}{r['rownum']}"), "values": [["Paid"]]})
        data.append({"range": safe_range(WL_TAB_NAME, f"{fee_col}{r['rownum']}"), "values": [[c["fee"]]]})
    try:
        ws.spreadsheet.values_batch_update({"valueInputOption": "USER_ENTERED", "data": data})
    except Exception as e:  # noqa: BLE001 — API 실패 시 어떤 행이 미기입인지 명시하고 종료
        rows_s = ", ".join(f"r{r['rownum']}" for r, _ in updates)
        raise SystemExit(f"[FATAL] 시트 기입 실패 ({e}) — 대상 행 전체 미반영 가능: {rows_s}. 재실행으로 복구 (멱등).")
    print(f"\n[APPLY] {len(updates)}행 기입 완료.")

    # 검증 리드백 — Deal Type + Fee 둘 다 확인
    _, rows_after = load_wl_tab(ws)
    after_by_rn = {r["rownum"]: r for r in rows_after}
    bad = []
    for r, c in updates:
        got = after_by_rn.get(r["rownum"], {})
        if got.get("deal", "").lower() != "paid" or fee_num(got.get("fee", "")) != c["fee"]:
            bad.append(f"r{r['rownum']} (기대 Paid/¥{c['fee']:,}, 실제 {got.get('deal', '?')}/{got.get('fee', '?')})")
    if bad:
        print(f"[VERIFY-FAIL] 리드백 불일치 {len(bad)}행: {'; '.join(bad)}")
        sys.exit(1)
    print(f"[VERIFY] 리드백 Deal Type=Paid + Fee 일치 확인: {len(updates)}/{len(updates)}행")


if __name__ == "__main__":
    main()
